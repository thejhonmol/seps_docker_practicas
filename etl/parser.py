"""Read cached XLSM values; never execute macros or recalculate source formulas."""
import hashlib
import io
import re
import zipfile
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl

GROUPS = ('CONAFIPS', 'FINANCOOP', 'Mutualistas', 'Segmento 1', 'Segmento 2', 'Segmento 3')
FIELDS = ('fecha_corte', 'segmento', 'entidad', 'es_total', 'tipo', 'codigo',
          'nombre', 'unidad', 'valor', 'estado', 'archivo', 'hoja', 'celda')


def tidy(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip()


def numeric(value):
    if isinstance(value, bool):
        raise ValueError('booleano no permitido')
    try:
        result = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError('valor no numérico') from exc
    if not result.is_finite():
        raise ValueError('valor no finito')
    return result


def layout(sheet):
    """Date header followed by entity header; forward-fill only dated columns."""
    preview = list(sheet.iter_rows(max_row=25))
    for row in preview:
        dates = [c for c in row if isinstance(c.value, (date, datetime))]
        if not dates:
            continue
        header_row = dates[0].row + 1
        headers = preview[header_row - 1]
        columns = []
        current = None
        for c, h in zip(row, headers):
            if isinstance(c.value, (date, datetime)):
                current = c.value.date() if isinstance(c.value, datetime) else c.value
            if current is not None and tidy(h.value):
                entity = tidy(h.value).upper()
                columns.append((c.column, current, entity, entity.startswith('VT_TOTAL')))
        if not columns:
            raise ValueError(f'{sheet.title}: no hay entidades bajo las fechas')
        return header_row, columns
    raise ValueError(f'{sheet.title}: encabezado de fechas no reconocido')


def parse_sheet(sheet, group, filename, kind):
    header, columns = layout(sheet)
    first_col = min(c[0] for c in columns)
    last_col = max(c[0] for c in columns)
    seen = set()
    for row in sheet.iter_rows(min_row=header + 1, max_col=last_col):
        if kind == 'cuenta':
            raw_code = row[0].value
            code = str(int(raw_code)) if isinstance(raw_code, (int, float)) and int(raw_code) == raw_code else tidy(raw_code)
            if not re.fullmatch(r'\d+', code):
                continue
            name = tidy(row[1].value)
            if not name:
                raise ValueError(f'{filename}/{sheet.title}: cuenta sin nombre')
        else:
            name = tidy(row[first_col - 2].value)
            if not name:
                continue
            cells = [row[col - 1] for col, *_ in columns]
            # Section labels have no values and no percentage formatting.
            if all(c.value is None and '%' not in c.number_format for c in cells):
                continue
            code = name.upper()
        for col, period, entity, total in columns:
            cell = row[col - 1]
            unit = 'USD' if kind == 'cuenta' else ('proporcion' if '%' in cell.number_format else 'numero')
            state, value = 'valido', None
            if cell.value is None:
                state = 'sin_dato'
            elif cell.data_type == 'e':
                state = 'error_excel'
            else:
                try:
                    value = numeric(cell.value)
                except ValueError:
                    state = 'no_numerico'
            key = (period, entity, kind, code)
            if key in seen:
                raise ValueError(f'Clave duplicada en origen: {group}/{key}')
            seen.add(key)
            record = dict(zip(FIELDS, (period.isoformat(), group, entity, total, kind,
                                      code, name, unit, value, state, filename,
                                      sheet.title, cell.coordinate)))
            if state != 'valido':
                record['motivo'] = str(cell.value) if cell.value is not None else 'Celda vacía o fórmula sin resultado guardado'
            yield record


def read_zip(path):
    """The uploaded six-workbook bundle; fail on unknown/missing/duplicate groups."""
    with zipfile.ZipFile(path) as archive:
        names = [n for n in archive.namelist() if n.lower().endswith('.xlsm')]
        found = Counter()
        for name in names:
            matches = [g for g in GROUPS if g.lower() in Path(name).name.lower()]
            if len(matches) != 1:
                raise ValueError(f'Boletín no reconocido: {name}')
            group = matches[0]
            found[group] += 1
            if found[group] > 1:
                raise ValueError(f'Boletín duplicado: {group}')
            workbook = openpyxl.load_workbook(io.BytesIO(archive.read(name)), read_only=True, data_only=True)
            try:
                for prefix, kind in [('2.', 'cuenta'), ('5.', 'indicador')]:
                    sheets = [s for s in workbook.worksheets if s.title.startswith(prefix)]
                    if len(sheets) != 1:
                        raise ValueError(f'{name}: hoja {prefix} ausente o ambigua')
                    yield from parse_sheet(sheets[0], group, Path(name).name, kind)
            finally:
                workbook.close()
        if set(found) != set(GROUPS):
            raise ValueError(f'Faltan boletines: {sorted(set(GROUPS)-set(found))}')


def sha256(path):
    with open(path, 'rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()
